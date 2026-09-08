from rest_framework import viewsets, status
from rest_framework.decorators import action, api_view
from rest_framework.response import Response
from django.db import models
from django.db.models import Q, Sum, Count
from .models import Deal, ContactDeal, Product, LineItem
from .serializers import (
    DealSerializer, ContactDealSerializer,
    ProductSerializer, LineItemSerializer,
)
from qontak_sales.apps.notifications.views import create_notification
from qontak_sales.apps.accounts.permissions import IsOwnerOrManager, IsManagerOrReadOnly


class DealViewSet(viewsets.ModelViewSet):
    serializer_class = DealSerializer
    permission_classes = [IsOwnerOrManager]

    def get_queryset(self):
        user = self.request.user
        qs = Deal.objects.filter(company__company=user.company, is_archived=False)
        if user.role == "AGENT":
            qs = qs.filter(owner=user)

        search = self.request.query_params.get("search")
        if search:
            qs = qs.filter(
                Q(name__icontains=search) |
                Q(company__name__icontains=search) |
                Q(description__icontains=search)
            )

        stage = self.request.query_params.get("stage")
        if stage:
            qs = qs.filter(stage=stage)

        company_id = self.request.query_params.get("company_id")
        if company_id:
            qs = qs.filter(company_id=company_id)

        owner = self.request.query_params.get("owner")
        if owner:
            qs = qs.filter(owner_id=owner)

        date_from = self.request.query_params.get("date_from")
        date_to = self.request.query_params.get("date_to")
        if date_from:
            qs = qs.filter(expected_close_date__gte=date_from)
        if date_to:
            qs = qs.filter(expected_close_date__lte=date_to)

        ordering = self.request.query_params.get("ordering", "-created_at")
        allowed_orderings = [
            "name", "-name", "amount", "-amount", "stage", "-stage",
            "expected_close_date", "-expected_close_date",
            "created_at", "-created_at", "probability", "-probability",
        ]
        if ordering in allowed_orderings:
            qs = qs.order_by(ordering)

        return qs

    def perform_create(self, serializer):
        deal = serializer.save(owner=self.request.user)
        # Notify team members (managers) about new deal
        from django.contrib.auth import get_user_model
        User = get_user_model()
        managers = User.objects.filter(company=self.request.user.company, role="MANAGER")
        for mgr in managers:
            if mgr != self.request.user:
                create_notification(mgr, "New Deal Created", f"Deal '{deal.name}' was created by {self.request.user.get_full_name() or self.request.user.username}", f"/deals/{deal.id}")

    @action(detail=True, methods=["post"])
    def move_stage(self, request, pk=None):
        deal = self.get_object()
        old_stage = deal.stage
        new_stage = request.data.get("stage")
        valid_stages = [s[0] for s in Deal.STAGE_CHOICES]
        if new_stage not in valid_stages:
            return Response({"error": "Invalid stage"}, status=400)

        deal.stage = new_stage
        if new_stage == "WON":
            deal.probability = 100
            from django.utils import timezone
            deal.actual_close_date = timezone.now().date()
        elif new_stage == "LOST":
            deal.probability = 0
            deal.lost_reason = request.data.get("lost_reason", "")
            deal.lost_notes = request.data.get("lost_notes", "")
            from django.utils import timezone
            deal.actual_close_date = timezone.now().date()
        elif new_stage == "QUALIFICATION":
            deal.probability = 10
        elif new_stage == "DISCOVERY":
            deal.probability = 30
        elif new_stage == "PROPOSAL":
            deal.probability = 50
        elif new_stage == "NEGOTIATION":
            deal.probability = 70
        elif new_stage == "CLOSING":
            deal.probability = 90

        deal.save()

        # Notify relevant people about stage change
        from django.contrib.auth import get_user_model
        User = get_user_model()
        managers = User.objects.filter(company=request.user.company, role="MANAGER")
        for mgr in managers:
            if mgr != request.user:
                create_notification(mgr, "Deal Stage Updated", f"Deal '{deal.name}' moved from {dict(Deal.STAGE_CHOICES).get(old_stage, old_stage)} to {dict(Deal.STAGE_CHOICES).get(new_stage, new_stage)}", f"/deals/{deal.id}")
        if deal.owner and deal.owner != request.user:
            create_notification(deal.owner, "Your Deal Stage Updated", f"Deal '{deal.name}' moved to {dict(Deal.STAGE_CHOICES).get(new_stage, new_stage)}", f"/deals/{deal.id}")

        return Response(DealSerializer(deal).data)

    @action(detail=True, methods=["post"])
    def archive(self, request, pk=None):
        deal = self.get_object()
        deal.is_archived = True
        deal.save()
        return Response({"message": "Deal archived"})

    @action(detail=True, methods=["post"])
    def restore(self, request, pk=None):
        deal = self.get_object()
        deal.is_archived = False
        deal.save()
        return Response({"message": "Deal restored"})

    @action(detail=False, methods=["get"])
    def pipeline(self, request):
        user = request.user
        deals = Deal.objects.filter(
            company__company=user.company, is_archived=False
        ).exclude(stage__in=["WON", "LOST"])
        if user.role == "AGENT":
            deals = deals.filter(owner=user)

        pipeline_data = {}
        for stage_code, stage_name in Deal.STAGE_CHOICES:
            if stage_code in ["WON", "LOST"]:
                continue
            stage_deals = deals.filter(stage=stage_code)
            pipeline_data[stage_code] = {
                "label": stage_name,
                "deals": DealSerializer(stage_deals, many=True).data,
                "count": stage_deals.count(),
                "total_value": stage_deals.aggregate(total=models.Sum("amount"))["total"] or 0,
            }

        won_deals = Deal.objects.filter(
            company__company=user.company, is_archived=False, stage="WON"
        )
        lost_deals = Deal.objects.filter(
            company__company=user.company, is_archived=False, stage="LOST"
        )
        if user.role == "AGENT":
            won_deals = won_deals.filter(owner=user)
            lost_deals = lost_deals.filter(owner=user)

        pipeline_data["WON"] = {
            "label": "Won",
            "deals": DealSerializer(won_deals, many=True).data,
            "count": won_deals.count(),
            "total_value": won_deals.aggregate(total=models.Sum("amount"))["total"] or 0,
        }
        pipeline_data["LOST"] = {
            "label": "Lost",
            "deals": DealSerializer(lost_deals, many=True).data,
            "count": lost_deals.count(),
            "total_value": lost_deals.aggregate(total=models.Sum("amount"))["total"] or 0,
        }

        return Response(pipeline_data)


class ContactDealViewSet(viewsets.ModelViewSet):
    serializer_class = ContactDealSerializer

    def get_queryset(self):
        user = self.request.user
        qs = ContactDeal.objects.filter(deal__company__company=user.company)
        deal_id = self.request.query_params.get("deal_id")
        if deal_id:
            qs = qs.filter(deal_id=deal_id)
        contact_id = self.request.query_params.get("contact_id")
        if contact_id:
            qs = qs.filter(contact_id=contact_id)
        return qs

    def perform_create(self, serializer):
        contact_deal = serializer.save()
        # Notify managers about new contact added to deal
        from django.contrib.auth import get_user_model
        User = get_user_model()
        managers = User.objects.filter(company=self.request.user.company, role="MANAGER")
        for mgr in managers:
            if mgr != self.request.user:
                create_notification(mgr, "Contact Added to Deal", f"{contact_deal.contact.full_name} added to deal '{contact_deal.deal.name}' by {self.request.user.get_full_name() or self.request.user.username}", f"/deals/{contact_deal.deal.id}")


class ProductViewSet(viewsets.ModelViewSet):
    serializer_class = ProductSerializer
    permission_classes = [IsManagerOrReadOnly]

    def get_queryset(self):
        user = self.request.user
        qs = Product.objects.filter(company=user.company)

        search = self.request.query_params.get("search")
        if search:
            qs = qs.filter(
                Q(name__icontains=search) |
                Q(code__icontains=search) |
                Q(category__icontains=search)
            )

        status_filter = self.request.query_params.get("status")
        if status_filter:
            qs = qs.filter(status=status_filter)

        category = self.request.query_params.get("category")
        if category:
            qs = qs.filter(category=category)

        return qs

    def perform_create(self, serializer):
        serializer.save(company=self.request.user.company)


class LineItemViewSet(viewsets.ModelViewSet):
    serializer_class = LineItemSerializer

    def get_queryset(self):
        user = self.request.user
        qs = LineItem.objects.filter(deal__company__company=user.company)
        deal_id = self.request.query_params.get("deal_id")
        if deal_id:
            qs = qs.filter(deal_id=deal_id)
        return qs

    def perform_create(self, serializer):
        line_item = serializer.save()
        deal = line_item.deal
        deal.amount = deal.line_items.aggregate(
            total=models.Sum("total_price")
        )["total"] or 0
        deal.save()
        # Notify managers about product added to deal
        from django.contrib.auth import get_user_model
        User = get_user_model()
        managers = User.objects.filter(company=self.request.user.company, role="MANAGER")
        for mgr in managers:
            if mgr != self.request.user:
                create_notification(mgr, "Product Added to Deal", f"{line_item.product.name} added to deal '{deal.name}' by {self.request.user.get_full_name() or self.request.user.username}", f"/deals/{deal.id}")

    def perform_update(self, serializer):
        serializer.save()
        deal = serializer.instance.deal
        deal.amount = deal.line_items.aggregate(
            total=models.Sum("total_price")
        )["total"] or 0
        deal.save()

    def perform_destroy(self, instance):
        deal = instance.deal
        instance.delete()
        deal.amount = deal.line_items.aggregate(
            total=models.Sum("total_price")
        )["total"] or 0
        deal.save()


@api_view(["GET"])
def dashboard_stats(request):
    from django.contrib.auth import get_user_model
    User = get_user_model()

    user = request.user
    all_deals = Deal.objects.filter(company__company=user.company, is_archived=False)
    if user.role == "AGENT":
        all_deals = all_deals.filter(owner=user)

    total_revenue = all_deals.filter(stage="WON").aggregate(total=Sum("amount"))["total"] or 0
    total_deals = all_deals.count()
    won_count = all_deals.filter(stage="WON").count()
    lost_count = all_deals.filter(stage="LOST").count()
    closed_count = won_count + lost_count
    win_rate = round((won_count / closed_count * 100) if closed_count > 0 else 0, 1)
    active_deals = all_deals.exclude(stage__in=["WON", "LOST"]).count()

    stage_dist = []
    for stage_code, stage_label in Deal.STAGE_CHOICES:
        count = all_deals.filter(stage=stage_code).count()
        stage_dist.append({"stage": stage_label, "count": count})

    monthly_data = []
    from django.db.models.functions import TruncMonth
    monthly = (
        all_deals.filter(stage="WON")
        .annotate(month=TruncMonth("created_at"))
        .values("month")
        .annotate(revenue=Sum("amount"), count=Count("id"))
        .order_by("month")[:12]
    )
    for m in monthly:
        monthly_data.append({
            "month": m["month"].strftime("%b %Y"),
            "revenue": float(m["revenue"]),
            "count": m["count"],
        })

    agents = User.objects.filter(company=user.company, role="AGENT")
    leaderboard = []
    for agent in agents:
        agent_won = all_deals.filter(owner=agent, stage="WON")
        agent_revenue = agent_won.aggregate(total=Sum("amount"))["total"] or 0
        leaderboard.append({
            "name": agent.get_full_name() or agent.username,
            "deals": agent_won.count(),
            "revenue": float(agent_revenue),
        })
    leaderboard.sort(key=lambda x: x["revenue"], reverse=True)

    return Response({
        "total_revenue": float(total_revenue),
        "win_rate": win_rate,
        "active_leads": active_deals,
        "total_leads": total_deals,
        "won_count": won_count,
        "lost_count": lost_count,
        "stage_distribution": stage_dist,
        "monthly_revenue": monthly_data,
        "leaderboard": leaderboard[:10],
    })


@api_view(["GET"])
def calendar_events(request):
    from django.utils import timezone
    from qontak_sales.apps.activities.models import ActivityLog

    user = request.user
    start = request.query_params.get("start")
    end = request.query_params.get("end")

    if not start or not end:
        return Response({"error": "start and end params required"}, status=400)

    deals = Deal.objects.filter(
        company__company=user.company, is_archived=False,
        expected_close_date__gte=start, expected_close_date__lte=end,
    )
    if user.role == "AGENT":
        deals = deals.filter(owner=user)

    deal_events = []
    for deal in deals:
        deal_events.append({
            "id": f"deal-{deal.id}",
            "type": "DEAL",
            "title": f"Deal: {deal.name}",
            "date": deal.expected_close_date.isoformat(),
            "time": "All Day",
            "deal_id": deal.id,
            "deal_name": deal.name,
            "color": "#2563EB",
        })

    activities = ActivityLog.objects.filter(
        deal__company__company=user.company,
        scheduled_at__isnull=False,
        scheduled_at__date__gte=start,
        scheduled_at__date__lte=end,
    )
    if user.role == "AGENT":
        activities = activities.filter(agent=user)

    activity_events = []
    TYPE_COLORS = {"CALL": "#2563EB", "EMAIL": "#8B5CF6", "MEETING": "#F59E0B", "NOTE": "#64748B", "FOLLOW_UP": "#059669"}
    for act in activities:
        local_sa = timezone.localtime(act.scheduled_at)
        activity_events.append({
            "id": f"activity-{act.id}",
            "type": act.activity_type,
            "title": f"{act.get_activity_type_display()}: {act.deal.name}",
            "date": local_sa.date().isoformat(),
            "time": local_sa.strftime("%I:%M %p"),
            "deal_id": act.deal_id,
            "deal_name": act.deal.name,
            "notes": act.notes,
            "is_completed": act.is_completed,
            "color": TYPE_COLORS.get(act.activity_type, "#64748B"),
        })

    return Response(deal_events + activity_events)


@api_view(["GET"])
def dashboard_export(request):
    import io
    from django.http import StreamingHttpResponse
    from django.utils import timezone
    from django.db.models.functions import TruncMonth
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    user = request.user
    all_deals = Deal.objects.filter(company__company=user.company, is_archived=False)
    if user.role == "AGENT":
        all_deals = all_deals.filter(owner=user)

    total_revenue = all_deals.filter(stage="WON").aggregate(total=Sum("amount"))["total"] or 0
    total_deals = all_deals.count()
    won_count = all_deals.filter(stage="WON").count()
    lost_count = all_deals.filter(stage="LOST").count()
    closed_count = won_count + lost_count
    win_rate = round((won_count / closed_count * 100) if closed_count > 0 else 0, 1)
    active_deals = all_deals.exclude(stage__in=["WON", "LOST"]).count()

    monthly = (
        all_deals.filter(stage="WON")
        .annotate(month=TruncMonth("created_at"))
        .values("month")
        .annotate(revenue=Sum("amount"), count=Count("id"))
        .order_by("month")[:12]
    )

    STAGE_MAP = dict(Deal.STAGE_CHOICES)

    primary_color = "2563EB"
    header_fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    title_font = Font(name="Calibri", bold=True, color=primary_color, size=16)
    subtitle_font = Font(name="Calibri", color="666666", size=10)
    section_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    section_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    label_font = Font(name="Calibri", bold=True, size=10)
    value_font = Font(name="Calibri", size=10)
    metric_value_font = Font(name="Calibri", bold=True, size=12, color=primary_color)
    zebra_light = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    def apply_border(ws, min_row, max_row, min_col, max_col):
        for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                cell.border = thin_border

    def apply_zebra(ws, min_row, max_row, min_col, max_col):
        for i, row in enumerate(ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col)):
            if i % 2 == 1:
                for cell in row:
                    cell.fill = zebra_light

    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard Report"
    ws.sheet_properties.tabColor = primary_color

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 22

    ws.merge_cells("A1:B1")
    ws["A1"] = "QontakSales Dashboard Report"
    ws["A1"].font = title_font
    ws["A1"].alignment = left_align

    ws.merge_cells("A2:B2")
    ws["A2"] = f"Generated: {timezone.now().strftime('%d %B %Y, %H:%M')}"
    ws["A2"].font = subtitle_font

    ws.merge_cells("A3:B3")
    ws["A3"] = f"Agent: {user.get_full_name() or user.username} ({user.role})"
    ws["A3"].font = subtitle_font

    row = 5
    ws.merge_cells(f"A{row}:B{row}")
    ws[f"A{row}"] = "SUMMARY"
    ws[f"A{row}"].font = section_font
    ws[f"A{row}"].fill = section_fill
    ws[f"A{row}"].alignment = center_align
    ws[f"B{row}"].fill = section_fill

    summary_items = [
        ("Total Revenue", f"Rp {total_revenue:,.0f}"),
        ("Win Rate", f"{win_rate}%"),
        ("Active Deals", str(active_deals)),
        ("Total Deals", str(total_deals)),
        ("Won Deals", str(won_count)),
        ("Lost Deals", str(lost_count)),
    ]

    for i, (label, val) in enumerate(summary_items):
        r = row + 1 + i
        ws[f"A{r}"] = label
        ws[f"A{r}"].font = label_font
        ws[f"A{r}"].alignment = left_align
        ws[f"B{r}"] = val
        ws[f"B{r}"].font = metric_value_font if i < 2 else value_font
        ws[f"B{r}"].alignment = right_align

    apply_border(ws, row, row + len(summary_items), 1, 2)
    apply_zebra(ws, row + 1, row + len(summary_items), 1, 2)

    monthly_start = row + len(summary_items) + 3
    ws.merge_cells(f"A{monthly_start}:B{monthly_start}")
    ws[f"A{monthly_start}"] = "MONTHLY REVENUE"
    ws[f"A{monthly_start}"].font = section_font
    ws[f"A{monthly_start}"].fill = section_fill
    ws[f"A{monthly_start}"].alignment = center_align
    ws[f"B{monthly_start}"].fill = section_fill

    mh_row = monthly_start + 1
    for col_idx, h in enumerate(["Month", "Revenue", "Deals Won"], 1):
        c = ws.cell(row=mh_row, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center_align

    mr = mh_row + 1
    for m in monthly:
        ws.cell(row=mr, column=1, value=m["month"].strftime("%b %Y")).font = value_font
        ws.cell(row=mr, column=1).alignment = left_align
        rev_cell = ws.cell(row=mr, column=2, value=float(m["revenue"]))
        rev_cell.font = value_font
        rev_cell.number_format = '"Rp "#,##0'
        rev_cell.alignment = right_align
        ws.cell(row=mr, column=3, value=m["count"]).font = value_font
        ws.cell(row=mr, column=3).alignment = center_align
        mr += 1

    if mr == mh_row + 1:
        ws.cell(row=mr, column=1, value="No data").font = Font(italic=True, color="999999")
        mr += 1

    apply_border(ws, monthly_start, mr - 1, 1, 3)
    apply_zebra(ws, mh_row + 1, mr - 1, 1, 3)

    deals_start_col = 4
    deals_header_row = 5
    deals_headers = ["Name", "Company", "Amount", "Stage", "Probability", "Expected Close", "Owner"]

    ws.merge_cells(f"D{deals_header_row}:G{deals_header_row}")
    ws[f"D{deals_header_row}"] = "DEALS DATA"
    ws[f"D{deals_header_row}"].font = section_font
    ws[f"D{deals_header_row}"].fill = section_fill
    ws[f"D{deals_header_row}"].alignment = center_align
    for ci in range(deals_start_col + 1, deals_start_col + len(deals_headers)):
        ws.cell(row=deals_header_row, column=ci).fill = section_fill

    dh_row = deals_header_row + 1
    for col_idx, h in enumerate(deals_headers, deals_start_col):
        c = ws.cell(row=dh_row, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center_align

    dr = dh_row + 1
    for deal in all_deals.select_related("owner", "company"):
        ws.cell(row=dr, column=4, value=deal.name).font = value_font
        ws.cell(row=dr, column=4).alignment = left_align
        ws.cell(row=dr, column=5, value=deal.company.name if deal.company else "").font = value_font
        ws.cell(row=dr, column=5).alignment = left_align
        val_cell = ws.cell(row=dr, column=6, value=float(deal.amount))
        val_cell.font = value_font
        val_cell.number_format = '"Rp "#,##0'
        val_cell.alignment = right_align
        ws.cell(row=dr, column=7, value=STAGE_MAP.get(deal.stage, deal.stage)).font = value_font
        ws.cell(row=dr, column=7).alignment = center_align
        ws.cell(row=dr, column=8, value=f"{deal.probability}%").font = value_font
        ws.cell(row=dr, column=8).alignment = center_align
        ws.cell(row=dr, column=9, value=deal.expected_close_date.strftime("%d %b %Y") if deal.expected_close_date else "").font = value_font
        ws.cell(row=dr, column=9).alignment = center_align
        ws.cell(row=dr, column=10, value=deal.owner.get_full_name() if deal.owner else "Unassigned").font = value_font
        ws.cell(row=dr, column=10).alignment = left_align
        dr += 1

    if dr == dh_row + 1:
        ws.cell(row=dr, column=4, value="No deals found").font = Font(italic=True, color="999999")
        dr += 1

    apply_border(ws, deals_header_row, dr - 1, 4, 10)
    apply_zebra(ws, dh_row + 1, dr - 1, 4, 10)

    ws.freeze_panes = "D6"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"dashboard-report-{timezone.now().strftime('%Y%m%d')}.xlsx"

    response = StreamingHttpResponse(
        output,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
